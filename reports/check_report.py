from pathlib import Path
from zipfile import ZipFile
import re
from lxml import etree
from docx import Document

root = Path(__file__).resolve().parent
path = root / 'AnonSpace_Final_Report.docx'
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(path) as z:
    assert z.testzip() is None
    xml = etree.fromstring(z.read('word/document.xml'))
    pages = ['']
    for node in xml.iter():
        local = etree.QName(node).localname
        if local == 'lastRenderedPageBreak':
            pages.append('')
        elif local == 't':
            pages[-1] += (node.text or '') + ' '
    for i, content in enumerate(pages, 1):
        print(f'{i:02} | {len(content.split()):3} words | {content[:110]} ... {content[-85:]}')
    assert not xml.xpath('.//w:drawing', namespaces=ns), 'Unexpected inserted diagram/image'
    text = ' '.join(xml.xpath('.//w:t/text()', namespaces=ns))
    for unwanted in ['AzureMed', 'Lotus Shrine', 'Oak Soe', 'MySQL', 'Vite', 'Error!']:
        assert unwanted not in text, f'Unexpected text: {unwanted}'
    names = set(xml.xpath('.//w:bookmarkStart/@w:name', namespaces=ns))
    instructions = [''.join(p.xpath('.//w:instrText/text()', namespaces=ns))
                    for p in xml.xpath('.//w:p', namespaces=ns)]
    targets = [re.search(r'PAGEREF\s+(\w+)', ins).group(1) for ins in instructions if 'PAGEREF' in ins]
    assert all(t in names for t in targets), 'Broken page reference'
    assert len(targets) == 99, f'Expected 99 index entries, got {len(targets)}'
    d = Document(path)
    headings = [p.text for p in d.paragraphs if p.style.name == 'Heading 1']
    assert len(headings) == 7
    placeholders = [t for t in d.tables if len(t.rows) == 1 and len(t.columns) == 1]
    assert len(placeholders) == 38
    print(f'PASS: valid Word package; {len(headings)} chapters; {len(placeholders)} figure spaces; {len(targets)} valid page references.')
