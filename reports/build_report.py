"""Create the AnonSpace school report with editable Word text and figure spaces."""
from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'AnonSpace_Final_Report.docx'
doc = Document()
doc.core_properties.title = 'AnonSpace — Final Project Report'
doc.core_properties.subject = 'Anonymous Social Networking and Community Platform'
doc.core_properties.author = 'AnonSpace Project Team'
doc.core_properties.keywords = 'AnonSpace, final report, student project, social network'

for s in doc.sections:
    s.page_width, s.page_height = Inches(8.268), Inches(11.693)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    s.header_distance = s.footer_distance = Inches(.5)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for name, size in [('Heading 1', 14), ('Heading 2', 14), ('Heading 3', 14)]:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(10)
    style.paragraph_format.keep_with_next = True
for name in ['Caption', 'Title', 'Subtitle']:
    doc.styles[name].font.name = 'Times New Roman'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Caption'].font.size = Pt(11)
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].paragraph_format.space_after = Pt(8)
doc.styles['Caption'].paragraph_format.line_spacing = 1.15

toc = []
figures = []
bookmark_id = 0
break_next = False

def prepare(para):
    global break_next
    if break_next:
        para.paragraph_format.page_break_before = True
        break_next = False
    return para

def bookmark(p, name):
    global bookmark_id
    bookmark_id += 1
    start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
    start.set(qn('w:id'), str(bookmark_id))
    start.set(qn('w:name'), name)
    end.set(qn('w:id'), str(bookmark_id))
    p._p.insert(0, start)
    p._p.append(end)

def field(p, instruction, fallback='1'):
    r = p.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    code = OxmlElement('w:instrText')
    code.set(qn('xml:space'), 'preserve')
    code.text = f' {instruction} '
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    value = OxmlElement('w:t')
    value.text = fallback
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for el in [begin, code, sep, value, end]:
        r._r.append(el)

def p(text, bold=False, center=False, size=None):
    para = prepare(doc.add_paragraph())
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    run = para.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return para

def body(text):
    for paragraph in text.strip().split('\n\n'):
        p(' '.join(paragraph.split()))

def label(title, text):
    para = prepare(doc.add_paragraph())
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.add_run(title + ' ').bold = True
    para.add_run(text)
    return para

def page():
    global break_next
    break_next = True

def heading(number, title, level=2):
    para = prepare(doc.add_paragraph(f'{number} {title}', style=f'Heading {level}'))
    name = 'sec_' + number.replace('.', '_')
    bookmark(para, name)
    toc.append((number, title, name, level))
    return para

def chapter(number, title, first=False):
    if not first:
        page()
    para = prepare(doc.add_paragraph(style='Heading 1'))
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(18)
    para.add_run(f'Chapter {number}\n').font.size = Pt(16)
    para.add_run(title).font.size = Pt(14)
    name = f'chapter_{number}'
    bookmark(para, name)
    toc.append((str(number), title, name, 1))

def new_section(fmt, start=1):
    s = doc.add_section(WD_SECTION_START.NEW_PAGE)
    s.header.is_linked_to_previous = False
    s.footer.is_linked_to_previous = False
    num = OxmlElement('w:pgNumType')
    num.set(qn('w:fmt'), fmt)
    num.set(qn('w:start'), str(start))
    s._sectPr.append(num)
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after = Pt(0)
    field(footer, 'PAGE')
    return s

def cell_text(cell, text, bold=False, size=10.5):
    para = cell.paragraphs[0]
    para.paragraph_format.line_spacing = 1.1
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.space_before = Pt(4)
    run = para.add_run(text)
    run.bold, run.font.size = bold, Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return para

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for c, width in zip(t.columns, widths):
            c.width = Inches(width)
    for c, title in zip(t.rows[0].cells, headers):
        cell_text(c, title, True)
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'F2F2F2')
        c._tc.get_or_add_tcPr().append(shd)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for c, value in zip(cells, row):
            cell_text(c, value)
    for row in t.rows:
        cant = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant)
        if widths:
            for c, width in zip(row.cells, widths):
                c.width = Inches(width)
    return t

def figure(number, title, height=5.2, kind='diagram'):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Inches(6.15)
    row = t.rows[0]
    row.height = Inches(height)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cant = OxmlElement('w:cantSplit')
    row._tr.get_or_add_trPr().append(cant)
    c = row.cells[0]
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders = OxmlElement('w:tcBorders')
    for edge in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement('w:' + edge)
        b.set(qn('w:val'), 'dashed')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), 'B7B7B7')
        borders.append(b)
    c._tc.get_or_add_tcPr().append(borders)
    cp = c.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.line_spacing = 1
    target = ' '.join(filter(None, [title.lower(), 'screenshot' if kind == 'UI' else kind]))
    r = cp.add_run(f'[Insert {target} here]')
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string('777777')
    cap = doc.add_paragraph(f'Figure {number}: {title}', style='Caption')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name = 'fig_' + number.replace('.', '_')
    bookmark(cap, name)
    figures.append((number, title, name))

# Cover: the reference's centred academic layout, without copying its identities.
p('[NAME OF UNIVERSITY / INSTITUTE]', True, True, 14)
p('[Department of Computer Science and Engineering]', False, True, 12)
logo = doc.add_paragraph('[Insert institution logo here]')
logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
logo.paragraph_format.space_before = Pt(30)
logo.paragraph_format.space_after = Pt(36)
logo.runs[0].font.color.rgb = RGBColor.from_string('777777')
p('Special Project', True, True, 14)
p('[Semester and Academic Year]', True, True, 14)
p('[Course Code and Project Level]', True, True, 14)
p('Anonymous Social Networking and Community Platform', True, True, 14)
brand = p('AnonSpace', True, True, 22)
brand.paragraph_format.space_before = Pt(20)
brand.paragraph_format.space_after = Pt(20)
p('Final Report', True, True, 16)
p('Submitted by', True, True)
p('[Student name(s) and roll number(s)]', False, True)
p('Supervised by', True, True)
p('[Supervisor name and designation]', False, True)
p('[Submission month and year]', False, True)

new_section('lowerRoman')
p('Acknowledgement', True, True, 14)
body('''We would like to thank [Name of University / Institute] and the Department of [Department Name] for giving us the opportunity to develop AnonSpace as our special project. This project gave us a chance to apply what we learned in web development, database management, and software engineering to a working application.

We would especially like to thank our project supervisor, [Supervisor Name], for the guidance and feedback provided during the project. The suggestions helped us think more carefully about the system requirements, the way users move between pages, and the responsibilities of administrators.

We are also grateful to our teachers, classmates, and the people who supported our work. Finally, we would like to acknowledge the effort of each member of our project team. Preparing this project helped us understand the importance of communication, checking each other's work, and making improvements step by step.''')

page()
p('Certificate', True, True, 14)
body('''This is to certify that the project report entitled “AnonSpace: Anonymous Social Networking and Community Platform” is submitted by [Student Name(s) and Roll Number(s)] of the Department of [Department Name], [Name of University / Institute], in partial fulfilment of the requirements of [Course / Special Project Name].

The report describes the design and implementation of the AnonSpace web application. The project is submitted for academic assessment under the supervision of [Supervisor Name and Designation].''')
p('[To be completed and signed by the project supervisor]', False, False, 11)
for _ in range(3):
    p('')
p('Project Supervisor', True)
p('Name: ____________________________________')
p('Designation: _______________________________')
p('Signature: __________________________________')
p('Date: ______________________________________')

page()
p('Abstract', True, True, 14)
body('''AnonSpace is a web-based social networking project that allows people to share posts and take part in discussions using automatically generated names. The idea behind the project is to give users a place where they can focus on what is being discussed without displaying their real names in the normal social interface. Users still register with an email address, so the system provides a public anonymous identity rather than complete anonymity from the service itself.

The application includes email verification, login, password recovery, text and media posts, comments, replies, voting, saved posts, communities, notifications, and direct messaging. The For You feed uses account activity and community interests to arrange posts, while the Popular feed uses voting scores. Users can join communities, manage their own contributions, and report unsuitable content. Dark mode, mobile layouts, and feed scroll restoration support everyday browsing.

The system has three account roles: user, admin, and superadmin. Admins review reports and moderate content, accounts, and communities. Superadmins have all admin features and can also register admins, review their moderation activity, and revoke admin access. Moderation records help the team review how these actions were performed.

We developed the frontend with Next.js, React, TypeScript, and Tailwind CSS. The backend uses PHP, and MongoDB stores the application data. PHPMailer delivers verification and recovery emails. Through this project, we connected user interface design with authentication, database relationships, and access control. The result is a student project that demonstrates how an anonymous community platform can combine social interaction with moderation and supervision.''')

page()
p('Table of Contents', True, True, 16)
toc_anchor = doc.add_paragraph()
page()
p('List of Figures', True, True, 16)
fig_anchor = doc.add_paragraph()

new_section('decimal')
chapter(1, 'Introduction', first=True)
heading('1.1', 'Introduction to AnonSpace')
body('''Social networking websites are commonly used to share ideas, ask questions, and communicate with other people. However, some users are uncomfortable discussing personal experiences or asking basic questions when their real names are visible. They may worry about how other people will judge them, especially when a discussion involves a sensitive topic.

We developed AnonSpace to explore a different way of taking part in online discussions. Instead of requiring a real name in the public interface, the application assigns each regular user a generated handle. Users can write posts, join communities, and reply to others through this handle. The aim is to make the conversation more important than the identity of the person posting it.

An anonymous platform also needs clear rules. Hiding a public name should not mean that harmful content cannot be reported or reviewed. For this reason, our project combines social features with voting, reporting, blocking, and moderation tools. An admin can review a report and take action, while a superadmin can also supervise admin accounts.

AnonSpace is a web application that can be opened in a browser on a computer or a mobile phone. We chose this approach so users can access the same features without installing a separate application. The project also gives us practical experience in connecting a modern frontend to a PHP backend and a MongoDB database.''')
page()
heading('1.2', 'Project Overview and Scope')
body('''The project is organised around three account roles. A regular user participates in the social platform. An admin works in a separate moderation panel. A superadmin uses the same moderation features and has additional responsibility for managing admins. Visitors can reach the welcome, login, registration, and password recovery pages, while the main social pages require a valid account session.''')
table(['Role', 'Scope of access'], [
    ('User', 'Read feeds; create posts; comment and reply; vote; save and share; join communities; use chat, notifications, and personal settings.'),
    ('Admin', 'View the dashboard; review reports; inspect posts, users, and communities; moderate violations; view ban logs; manage own settings.'),
    ('Superadmin', 'Use every admin feature, register admins, inspect individual admin activity, and revoke admin access.')
], [.95, 5.3])
body('''The social features include text posts, photo collections, video posts, nested replies, upvotes and downvotes, saved posts, reposting, and sharing links. Communities organise discussions by topic. Users must join a user-created community before publishing in it. The general Public destination is available without joining a separate community.

The project also includes email OTP verification, password recovery, notifications, private conversations between users, and account settings. Community creators can update community details, manage members, and pin posts. The admin panel provides oversight of the platform, while the superadmin tools provide oversight of staff accounts.

Our scope is an anonymous community website for academic demonstration. It does not include online payments, voice calls, video calls, or a native mobile application. A generated name hides a real name from normal public activity, but the backend still stores account email addresses and internal identifiers. The project should therefore be understood as a platform with pseudonymous participation, not a service that guarantees complete anonymity.''')

chapter(2, 'Project Objectives and Our Schedules')
heading('2.1', 'Project Objectives')
body('''The main objective of AnonSpace is to build a usable social platform where people can participate through generated names while administrators can manage inappropriate activity. We divided this objective into the following practical goals.''')
for title, text in [
    ('Provide anonymous public identities.', 'Generate a unique handle for each regular account and display it with posts and comments. Allow a name refresh after the first six months and once in each following six-month account-age period.'),
    ('Implement account verification.', 'Use email OTP verification during registration and password recovery. Add an expiry time and resend cooldown so the verification process is controlled.'),
    ('Support everyday social interaction.', 'Let users publish text, photos, and videos, add comments and replies, vote on content, and save or share posts.'),
    ('Organise discussions through communities.', 'Allow users to create and join communities, view community rules, and publish only where they have the required membership.'),
    ('Provide useful feeds.', 'Offer a For You feed based on simple interest signals, a Popular feed based on voting scores, and a Saved page for bookmarked posts.'),
    ('Improve communication and feedback.', 'Include direct messaging, notifications about comments and moderation, and short confirmation messages after common actions.'),
    ('Provide moderation tools.', 'Allow users to report content and give admins the tools to review reports, inspect activity, and apply appropriate bans.'),
    ('Provide supervision of admins.', 'Give superadmins all moderation features and separate tools for registering admins, reviewing their actions, and revoking their access.'),
    ('Make the interface practical on mobile.', 'Use responsive layouts, accessible action buttons, a mobile navigation menu, and theme preferences that remain after a reload.')
]:
    label(title, text)
page()
heading('2.2', 'Project Schedule')
body('''The following table presents an indicative fourteen-week organisation of the project work. The week ranges are a planning structure rather than a verified record of our actual start and finish dates. They can be aligned with the academic project calendar when the final submission details are entered.''')
table(['Weeks', 'Activity', 'Description'], [
    ('1–2', 'Planning and analysis', 'Identify the problem, define user roles, agree on the scope, and list the main social and moderation requirements.'),
    ('3–4', 'System design', 'Plan page layouts, database collections, access rules, use cases, and application flowcharts.'),
    ('5–6', 'Development Phase I', 'Set up the frontend and backend; implement registration, email OTP, login, and generated names.'),
    ('7–8', 'Social feature integration', 'Connect posts, media uploads, comments, replies, voting, communities, and saved posts.'),
    ('9–10', 'Administration and communication', 'Develop report review, ban records, admin supervision, notifications, chat, and sharing.'),
    ('11', 'Usability improvements', 'Refine feed behaviour, mobile layouts, dark mode, confirmation messages, and navigation.'),
    ('12–13', 'Testing and correction', 'Check permissions and important user flows; address errors; review validation and deployment settings.'),
    ('14', 'Final documentation and presentation', 'Prepare the report, diagrams, screenshots, demonstration, and project presentation.')
], [.62, 1.62, 4.01])
body('''The stages depend on one another. For example, comments and saved posts need working account sessions and post records first. Report review depends on having content to inspect, and superadmin supervision depends on recording admin actions. Planning the work in this order makes it easier to integrate the features gradually.''')

chapter(3, 'Background and Core Technologies')
body('''AnonSpace separates the interface, application logic, and persistent data into three main parts. The interface is built with Next.js and React. PHP receives API requests and applies the application rules. MongoDB stores the account, discussion, and moderation records. This separation helps us understand where a problem occurs and which part needs to be changed.''')
heading('3.1', 'Frontend Development')
label('Next.js and React.', 'The frontend package uses Next.js 16 and React 19. Pages are organised with the App Router, and reusable React components handle post cards, comment threads, navigation, chat, and admin controls. Shared layouts give related pages a consistent structure.')
label('TypeScript.', 'TypeScript describes the expected shape of users, posts, comments, reports, and API responses. These definitions help us find mismatches while developing the application, such as using a field that is not included in the response.')
label('HTML and Tailwind CSS.', 'React components produce the page structure, while Tailwind CSS classes control spacing, colour, borders, and responsive behaviour. The same components can change layout at different screen widths, including the mobile admin navigation drawer.')
label('Icons and media.', 'Lucide React provides icons for common actions such as saving, reporting, opening menus, and changing the theme. Photos are displayed in post media components, while videos use a player that responds to whether the video is visible on screen.')
label('Browser interaction.', 'The Fetch API sends requests to the backend. Browser local storage keeps the theme preference, and IntersectionObserver helps control feed video playback. React state handles form inputs, loading indicators, and the results shown on each page.')
body('''We reused components where the same behaviour appeared in several places. For example, a post can appear in the main feed, a community, or a saved list. Keeping its actions in a shared component reduces repeated work and helps these pages behave consistently.''')
page()
heading('3.2', 'Backend Development')
body('''The backend is written in PHP and requires PHP 8.1 or later. The entry point receives an HTTP request and routes it to an action class. Each action is responsible for a particular operation, such as creating a comment, requesting an OTP, listing posts, or reviewing a report. Shared helper classes handle authentication, uploads, pagination, and moderation.''')
label('API responses.', 'Most requests and responses use JSON. File uploads use multipart form data. The backend validates input, checks the current account, performs the operation, and returns a result or a clear error response.')
label('Authentication and authorisation.', 'Passwords are stored as bcrypt hashes. After a successful login, the backend creates a signed JWT session and sets an HttpOnly cookie. Protected operations reload the current account from the database, so the live account role and banned status are checked on subsequent requests.')
label('Email delivery.', 'PHPMailer sends verification and password recovery emails through SMTP. OTP records contain a hash of the code, an expiry time, and an attempt count. A new code replaces the earlier code for that email address.')
label('Social operations.', 'Separate actions create posts and comments, record votes, save posts, join communities, and manage messages. Membership and ownership rules are checked by the relevant backend actions, including the rule that users can delete only their own comments.')
label('Moderation operations.', 'Shared moderation code changes content status, records the reason and actor, and creates a notification. Report-review actions and direct moderation actions use these common operations to keep the result consistent.')
label('Dependencies and configuration.', 'Composer manages PHP packages, including the MongoDB library and PHPMailer. Database connection details, email credentials, and the JWT signing secret are supplied through environment configuration rather than being part of the report or user interface.')
page()
heading('3.3', 'Database Management')
body('''MongoDB is the database used by AnonSpace. It stores documents in collections instead of storing rows in relational tables. Records still have clear relationships: a comment refers to a post, a membership connects a user and a community, and a message belongs to a chat thread. Object identifiers are used for many of these links.

The database setup script defines document validation rules and indexes. Unique indexes are used for account emails, handles, and community slugs. Additional indexes support common lookups, such as finding comments for a post. OTP records have an expiry index, while the verification code also checks expiry directly.''')
table(['Collection', 'Main information stored'], [
    ('users', 'Email, password hash, handle, role, account status, creation date, and notification preferences.'),
    ('otp_codes', 'Email verification code hash, attempt count, issue time, and expiry time.'),
    ('communities', 'Slug, name, description, visibility, creator, rules, images, and status.'),
    ('community_members', 'The connection between a user and a joined community.'),
    ('posts', 'Author, community, text, media links, voting counts, status, and optional repost reference.'),
    ('comments', 'Post reference, optional parent comment, author, text, vote counts, and status.'),
    ('votes / saved_posts', 'A user’s vote on a post or comment, and a user’s saved-post references.'),
    ('chat_threads / chat_messages', 'Conversation participants, reading state, messages, and attached media links.'),
    ('reports / ban_logs', 'Reported targets, review decisions, moderation reasons, responsible admins, and dates.'),
    ('blocked_users / notifications', 'Blocking relationships and account-specific notification records.'),
    ('admin_logs', 'Admin registration and revocation actions performed by a superadmin.')
], [1.62, 4.63])
body('''Photos and videos are stored as uploaded files, while their locations are recorded in the database. This makes persistent upload storage as important as persistent database storage. The application also keeps banned content records for moderation purposes. Banning a post therefore changes its visibility status rather than automatically deleting its database record.''')

chapter(4, 'Project Overview and Implementation')
body('''This chapter describes how the application parts work together and provides spaces for the main system diagrams. The design connects the three account roles to the same backend, with different operations available according to each role.''')
heading('4.1', 'Overall Project Structure and Flow')
label('Frontend structure.', 'The frontend directory contains the application pages, shared components, reusable hooks, and the API client. Social pages share the main application layout. Admin and superadmin pages use the admin layout and sidebar.')
label('Backend structure.', 'The backend directory contains the PHP entry point, router, action classes, and shared service classes. An action checks its inputs and access requirements before it reads or changes database records.')
label('Database structure.', 'MongoDB collections store accounts, posts, comments, memberships, conversations, and moderation records. Backend helpers provide access to these collections and prepare response data for the frontend.')
label('Supporting services.', 'SMTP is used to send account verification emails. Uploaded media is served from the backend upload storage. Deployment configuration connects the browser frontend to the correct API address.')
body('''A typical request begins when the user performs an action on a page. For example, pressing the Comment button sends the post identifier and comment text to the backend. The backend verifies the session, confirms that the post is accessible, validates the comment, stores it, and updates the post’s comment count. It can also create a notification for the post author. The frontend then displays the returned comment.

The same pattern applies to administrative actions, but the permission check is different. Reviewing a report requires an admin or superadmin account. Registering an admin requires a superadmin account. These backend checks are important because simply hiding a button in the interface does not protect the underlying operation.''')
page()
heading('4.2', 'Use Case Diagram')
body('''The use case diagram should show the user, admin, and superadmin actors. User actions include registration, login, publishing, commenting, joining communities, messaging, and reporting. Admin actions include reviewing reports and moderating content. The superadmin actor inherits all admin use cases and adds registering admins, viewing admin activity, and revoking admin access.''')
figure('4.2', 'Use Case Diagram', kind='')
page()
heading('4.3', 'Application Flowcharts')
heading('4.3.1', 'Main Application Flowchart', 3)
body('''The main flow starts at the welcome page. A visitor registers or logs in. After authentication, a regular user enters the newsfeed, an admin enters the admin dashboard, and a superadmin enters the Admins section. Each role can then navigate to its available features. Logging out clears the session and returns the user to the entry page.''')
figure('4.3.1', 'Main Application Flowchart', kind='')
page()
heading('4.3.2', 'Community and Communication Services Flowchart', 3)
body('''This flowchart covers community and communication services. A user can discover a community, view its details, join it, and open its discussions. Community creators can edit details and manage members. A user can also open a conversation or notification. Before restricted content is returned, the backend checks the account and the relevant access rules.''')
figure('4.3.2', 'Community and Communication Services Flowchart', kind='')
page()
heading('4.3.3', 'Post and Comment Flowchart', 3)
body('''This flow begins when a user chooses a destination and prepares a post. The backend checks the community and membership, validates the text and media, and saves the post. Readers can vote, comment, reply, save, report, or share it. An author may delete their own comment; a banned comment is displayed as “Banned comment” so the reply chain can remain visible.''')
figure('4.3.3', 'Post and Comment Flowchart', kind='')
page()
heading('4.3.4', 'Administration and Supervision Flowchart', 3)
body('''An admin opens a report or moderation target, reviews the information, and chooses whether to confirm a violation or dismiss the report. A ban updates the target and creates a moderation record. A superadmin can follow this same flow and can additionally open the Admins section to register an admin, inspect that admin’s activity, or revoke admin access.''')
figure('4.3.4', 'Administration and Supervision Flowchart', kind='')
page()
heading('4.3.5', 'Account Management Flowchart', 3)
body('''Registration begins with an email address and OTP delivery. A valid code produces a verification ticket, after which the user sets a password and receives a generated handle. Invalid or expired codes require correction or a new request, subject to the resend cooldown. Existing users can log in, change their password, recover a forgotten password, or refresh their handle when the six-month rule allows it.''')
figure('4.3.5', 'Account Management Flowchart', kind='')
page()
heading('4.4', 'Entity-Relationship (ER) Diagram')
body('''Although AnonSpace uses MongoDB, an ER diagram can still show the logical relationships between its records. A user authors many posts and comments. A post has many comments, and a comment may have child replies. Memberships connect users and communities, saved records connect users and posts, and chat threads contain messages. Reports and logs refer to moderation targets and the accounts responsible for the actions.''')
figure('4.4', 'Entity-Relationship (ER) Diagram', kind='')

chapter(5, 'Implementation and Features')
heading('5.1', 'Key Feature Implementation')
label('Account registration and login.', 'A regular user registers with an email address, verifies the email, and sets a password. The backend checks the verification ticket before creating the account. A generated handle becomes the public name. Login verifies the password hash and issues the session cookie used for protected requests.')
label('OTP verification and recovery.', 'Verification codes contain six digits and expire after ten minutes. The OTP helper limits verification attempts and applies a thirty-second resend cooldown. Registration and recovery pages display the waiting time before another code can be requested. Email messages use an HTML layout, with a plain-text alternative for mail clients that need it.')
label('Anonymous name refresh.', 'The first name refresh becomes available six months after account creation. Later opportunities follow the six-, twelve-, eighteen-month pattern from the same creation date. A user can refresh once during an eligible period. Waiting before using an opportunity does not shift the entire schedule forward. The user keeps the same account identifier after changing the handle.')
label('Posts and media.', 'Users publish text in the Public destination or in a community they have joined. A post can include photos or a video. The backend currently accepts up to ten photos, with an eight-megabyte limit per photo, or a video up to fifty megabytes. It validates media types and prevents a single post from combining photos and video.')
label('Community membership.', 'A user-created community has a name, slug, topic, description, visibility, and rules. Membership is stored separately from the community. Posting requires membership, while the Public destination is treated as a general space. Private community content is restricted according to membership, and reposting cannot move private content into a different community.')
page()
label('Comments, replies, and deletion.', 'Comments belong to a post, and replies also store the identifier of their parent comment. Separate Comment and Reply buttons make submission clear. A user can delete only their own comment. When it has replies, those replies are moved up one level so other users’ contributions are retained. A comment that has already been banned is kept as a moderation placeholder.')
label('Preserving banned discussion chains.', 'Banning a comment sets its status to banned rather than deleting the record. The API replaces the visible text with “Banned comment”. Its relationship to the surrounding replies stays in place. This allows the thread to remain understandable while the original content is hidden from the normal comment response.')
label('Voting and automatic moderation.', 'Users can upvote or downvote posts and comments. A shared rule evaluates the downvote share after a minimum number of votes. In the inspected project version, the share is fifty percent and the minimum is four votes, explicitly marked in the code as a temporary testing setting. This threshold needs review before wider use because a small group could otherwise trigger a ban too easily.')
label('For You, Popular, and Saved.', 'The For You feed uses a scoring method based on joined communities, votes, saved posts, comments, freshness, and engagement. It is a rule-based ranking method rather than a trained machine-learning model. Popular orders posts by upvotes minus downvotes, with newer posts used to break ties. Saved shows the posts bookmarked by the current user.')
label('Saving and sharing.', 'Saving a post updates the saved-post record and displays a small confirmation for three seconds. The share menu supports copying a link and preparing it for a chat conversation. Reposting creates a new post with a reference to the original post, so the original can be resolved when the repost is displayed.')
page()
label('Feed continuity and video playback.', 'The frontend keeps feed items, pagination state, and scroll positions in memory for returning to For You, Popular, and Saved during client-side navigation. This helps when a user opens a comment thread or settings and then returns. The cache is not a permanent history across a full browser restart. Feed videos attempt muted playback when at least sixty percent of the player is visible and pause when it moves out of view; manual controls remain available.')
label('Notifications.', 'When another user comments on a post, the backend creates a notification for its author. It also creates notifications for report outcomes and moderation events. Users can open notifications, mark them as read, and control supported notification categories in settings. The notification indicator refreshes periodically, so delivery in the interface is not an instant push service.')
label('Direct messaging and blocking.', 'Chat threads connect participants and store individual messages, attached media, and reading information. A shared frontend context manages the desktop chat panel and the mobile chat route. Conversation updates are requested at intervals. Blocking records unwanted account relationships and is checked in supported interaction paths.')
label('Admin moderation.', 'Admins use a separate dashboard to view statistics and recent moderation activity. They can review reports about posts, comments, users, and communities, inspect content, and take moderation actions. Report decisions and bans record the responsible admin and time, allowing the action history to be reviewed later.')
label('Superadmin supervision and registration.', 'A superadmin has every admin navigation item and moderation capability. The additional Admins section lists staff accounts, provides registration, displays individual moderation histories, and supports revocation. Admin registration follows an email verification process and records the person’s name. A normal admin cannot use these staff-management operations because they require a superadmin role in the backend.')
label('Theme and account settings.', 'Dark mode is the default when no preference has been stored. Choosing light or dark mode writes the preference to local storage, and the page applies it early on the next visit. Account settings also provide password changes, notification preferences, and eligible name refresh actions.')
page()
heading('5.2', 'Component Communication and State Management')
body('''React components communicate through properties and callbacks. A parent page provides data to a post card or comment component, and the child reports a completed action so the parent can update the screen. Local state keeps track of form values, expanded replies, loading indicators, and errors.

The API client is collected in one module instead of repeating request logic in every page. It includes credentials when calling protected endpoints and turns failed responses into application errors. Pages can then show a suitable message, such as an expired session or a request that the current role is not allowed to perform.

Shared chat state is managed with React Context. This lets a post share action open the conversation picker and pass a draft link without every component having to manage a separate chat window. The theme uses browser storage and updates the root page class so that all components follow the chosen appearance.

Feed state is kept separately for each feed. Keeping the loaded posts as well as the scroll position matters: restoring only the scroll position would not work correctly if the page had returned to a shorter list. The For You feed also keeps its cursor, while paginated feeds keep the next page and whether more results are available.

The database remains the source of persistent application data. Browser state makes navigation smoother, but the backend still decides whether a requested action is allowed. A saved visual state is therefore not a substitute for checking the current account, content visibility, or community membership.''')
heading('5.3', 'Build and Deployment Considerations')
body('''The project contains separate frontend and backend applications. Both need their dependencies and configuration before the complete site can run. The following points describe the deployment arrangement documented in the project; they do not claim that a particular public deployment has been checked for this report.''')
label('Frontend build.', 'The frontend uses npm to install dependencies and Next.js to create the production build. The normal development command is npm run dev, and npm run build prepares the production application. ESLint and TypeScript checks help detect source-level problems before deployment.')
page()
label('Backend environment.', 'The PHP service needs Composer dependencies, the MongoDB extension, a database connection, SMTP access, and persistent storage for uploads. The included Dockerfile supports a PHP and Apache deployment. The project deployment guide describes Railway as a backend hosting option and Vercel for the frontend.')
label('API and cookie configuration.', 'The frontend must use the correct API base URL. The current route guard relies on receiving the session cookie, so the documented production arrangement places the frontend and backend on subdomains of the same parent domain. Cookie domain, allowed origin, HTTPS, and production cookie settings must agree with that arrangement.')
label('Storage and email.', 'Uploaded photos and videos need a persistent volume or equivalent storage so they survive a redeployment. Database and mail credentials belong in the hosting environment. SMTP delivery also needs to be checked with the configured sender, since a successful local form submission alone does not prove that an email reached an inbox.')
label('Validation and error handling.', 'The backend validates identifiers, text lengths, media formats, and access conditions before performing the relevant operation. Frontend loading and error states give feedback when an action is still pending or fails. These controls reduce common mistakes but do not replace a full security review.')
page()
label('Verification approach.', 'Important checks should cover the full path from a user action to its stored result. The table below records the expected outcomes for project verification. It is a checklist, not a claim that every scenario has already passed a recorded end-to-end test.')
table(['Scenario', 'Expected outcome'], [
    ('OTP resend and expiry', 'A resend is blocked during the thirty-second wait; an expired or invalid code cannot complete verification.'),
    ('Community posting', 'A non-member cannot publish in a user-created community; a joined member can submit valid content.'),
    ('Comment ownership', 'The author can delete their own comment; another user cannot delete it; existing replies are retained.'),
    ('Banned comment', 'The text becomes “Banned comment” in the thread and the reply relationships remain.'),
    ('Role permissions', 'Admin and superadmin can moderate; only superadmin can register or revoke admins.'),
    ('Saved feed and navigation', 'Saving gives feedback; returning through client navigation restores the loaded feed and scroll position.'),
    ('Theme and mobile layout', 'The selected theme survives reload; primary navigation and actions are usable on narrow screens.'),
    ('Messaging and notifications', 'Authorised participants can access a conversation; relevant notification records appear after refresh.')
], [1.6, 4.65])

chapter(6, 'User Interface and Experience')
heading('6.1', 'Overview')
body('''The AnonSpace interface is designed around reading and responding to discussions. For a regular user, the newsfeed is the main starting point. Posts display the author’s generated handle, community, content, and action buttons. Related pages provide communities, saved posts, messages, notifications, and settings.

Admins work in a separate interface because their main tasks are reviewing activity and making moderation decisions. The superadmin uses this same interface with an additional Admins section. This keeps shared moderation tasks familiar while making staff supervision easy to find.''')
heading('6.2', 'Design Principles')
for title, text in [
    ('Consistency.', 'Shared navigation, post cards, buttons, and form styles make related pages behave similarly. The same admin features are available to both admin and superadmin accounts.'),
    ('Simplicity.', 'The main feed focuses on the content and the actions that readers need. Additional options are grouped in menus, while forms show the information needed for the current step.'),
    ('Readable appearance.', 'The default dark theme uses dark slate backgrounds, light text, and cyan accents. A light theme is also available. Labels and icons help users recognise actions without depending only on colour.'),
    ('Responsiveness.', 'Layouts adjust for narrow screens. The admin sidebar becomes a drawer, while wide admin tables can scroll within their own area. Mobile users can reach the main social pages through compact navigation.'),
    ('Clear feedback.', 'Loading states, validation messages, confirmation prompts, unread indicators, and short save notifications tell users what happened after an action.'),
    ('Continuity.', 'Remembering the selected theme and restoring feed position during navigation help users continue from where they left off.')
]:
    label(title, text)
page()
body('''The colour palette below can be added from the final interface. The main colours should show the dark background, card surface, primary text, muted text, cyan action colour, and the warning or error colour. Screenshots in the next section should use the same final version of the application.''')
figure('6.1', 'Color Palette of AnonSpace', height=4.2, kind='image')
body('''The following pages leave space for the final UI screenshots. Each caption identifies the page or function that should be shown. These spaces are intentionally editable so that desktop and mobile captures can be inserted after the final interface is ready.''')

UI = [
    ('Home Page', 'The welcome page introduces AnonSpace and gives visitors clear entry points for login and registration. It is the first screen before entering the protected social pages.'),
    ('Login Page', 'Existing users enter their account credentials here. After a successful login, the application opens the destination for the user’s role.'),
    ('Sign-Up Page', 'The registration page begins with email verification and then collects the password. A regular account receives a generated public handle.'),
    ('OTP Verification Page', 'The verification screen provides the code input, validation feedback, and the resend countdown. It should show the thirty-second waiting state after resending.'),
    ('Password Recovery Page', 'A user who forgets a password can request email verification and then set a replacement password through the recovery flow.'),
    ('For You Newsfeed Page', 'This page shows the personalised order of posts with voting, comments, saving, and sharing actions. Returning during client navigation restores its previous feed state.'),
    ('Popular Newsfeed Page', 'The Popular feed displays posts ordered by voting score. It offers the same post interactions and keeps its browsing state separate from For You.'),
    ('Saved Posts Page', 'This page collects the posts saved by the current user. A saved post can be opened again or removed from the saved list.'),
    ('Create Post Page', 'The post form provides a destination, text input, and media selection. Publishing in a user-created community requires membership.'),
    ('Post Detail and Comment Page', 'The detail page shows the complete post and its discussion. Comment and Reply buttons make submission clear, and nested replies show the conversation structure.'),
    ('Share and Repost Interface', 'The share interface provides copying a post link and sending it through chat. Reposting allows the user to select an eligible destination for the original post reference.'),
    ('Communities Page', 'The communities area helps users discover groups and access communities they have joined. Community details explain the subject of each group.'),
    ('Community Detail Page', 'A community page displays its description, rules, membership state, and posts. The available posting and management actions depend on membership and ownership.'),
    ('Create Community Page', 'This form collects the information required for a new community, including its name, topic, description, and visibility.'),
    ('Community Management Interface', 'Community creators can edit details, manage members, and organise highlighted content. This screenshot should show the creator’s available controls.'),
    ('Search Page', 'Search provides a way to find relevant content and communities. Results should make it clear which item will open when selected.'),
    ('Profile Page', 'The user’s own profile provides access to their contributions and account-related information. Public participation continues to use the generated handle.'),
    ('Notifications Page', 'The notification page lists comments on the user’s posts and supported moderation events. Read and unread states help the user identify new activity.'),
    ('Chat Page', 'The chat interface shows conversation history and a message composer. On desktop, chat can also appear in a docked panel; mobile uses the dedicated chat page.'),
    ('Settings Page', 'Settings groups the theme, password, notification preferences, and name-refresh controls. The handle action is available only when the account-age rule allows it.'),
    ('Admin Dashboard Main Page', 'The dashboard summarises account and content activity, pending reports, communities, and moderation actions. Filters allow the admin to review relevant periods and categories.'),
    ('Admin Reports Page', 'Reports are organised by posts, accounts, communities, and comments. Admins review the reported target and either confirm a violation or dismiss the report.'),
    ('Admin Posts Page', 'The post-review page allows staff to inspect content and open detailed discussion views before deciding whether moderation is needed.'),
    ('Admin Users Page', 'The account lookup pages show user handles, account roles, status, and related activity for moderation. Staff accounts follow the separate superadmin management process.'),
    ('Admin Communities Page', 'This section supports looking up communities and inspecting their details and content. Moderation controls are available according to the staff role.'),
    ('Admin Ban Log Page', 'The ban log records moderation targets, reasons, and related action details. It helps staff understand whether a ban came from automatic rules or an admin decision.'),
    ('Superadmin Admins Page', 'The Admins page lists staff accounts and provides registration and revocation controls. Superadmins also retain every normal admin navigation item.'),
    ('Superadmin Admin Activity Page', 'Opening an admin account shows its moderation history. This supports supervision of bans and report decisions made by individual admins.'),
    ('Superadmin Register Admin Page', 'The registration flow collects the required admin information and verifies the email. The operation is available only to a superadmin.'),
    ('Mobile Admin and Superadmin Navigation', 'The mobile drawer makes moderation pages accessible on smaller screens. A superadmin sees the same admin links together with the additional Admins entry.')
]
for i, (title, description) in enumerate(UI, start=1):
    if i % 2 == 1:
        page()
    if i == 1:
        heading('6.3', 'User Interface Design')
    heading(f'6.3.{i}', title, 3)
    para = p(description)
    para.paragraph_format.line_spacing = 1.15
    para.paragraph_format.space_after = Pt(7)
    figure(f'6.{i+1}', title, height=2.5 if i > 2 else 2.25, kind='UI')

chapter(7, 'Conclusion and Future Enhancements')
heading('7.1', 'Conclusion')
body('''AnonSpace brings together the main parts of an anonymous community website in one web application. Regular users can register, participate through generated handles, publish posts, join communities, reply to discussions, vote, save content, and communicate through chat. Notifications and account settings support these daily activities.

An important part of the project is the connection between user freedom and moderation. Users can report unsuitable activity, admins can review it, and superadmins can supervise the admin accounts. Superadmin access includes the full admin feature set, with registration, activity review, and revocation added on top. Keeping these permissions clear helps us understand how the responsibilities fit together.

The project also shows why small details matter in a social application. A comment should not disappear together with other people’s replies just because its content is banned. A user should not be able to publish inside a community they have not joined. Theme settings should remain after a reload, and returning from a discussion should not unnecessarily lose the feed position. These requirements connect interface behaviour to backend rules and data structure.

Developing AnonSpace helped us practise component-based interface design, API communication, document database relationships, authentication, and role checks. It also helped us see the limits of the current version. A generated handle does not make the stored account data anonymous, periodic requests are different from instant messaging delivery, and a simple vote-based ban rule needs careful evaluation.

The current project provides a foundation for demonstrating these ideas in a school setting. Further user testing, deployment checks, and security review would help us judge how well it performs outside a controlled demonstration. The next improvements should build on these findings rather than assuming that every situation has already been covered.''')
page()
heading('7.2', 'Future Enhancements')
body('''The following improvements could be developed after the current project. They are proposed extensions and should not be understood as completed features.''')
for title, text in [
    ('More immediate communication.', 'Replace periodic chat and notification requests with a suitable push connection, such as WebSockets or server-sent events, and add optional browser notifications.'),
    ('Stronger moderation safeguards.', 'Review the temporary auto-ban threshold using realistic usage data. Add safeguards against coordinated downvoting and provide an appeal process so mistaken actions can be reviewed.'),
    ('Improved privacy controls.', 'Add clearer account-data explanations, data export, and account deletion workflows. Review uploaded media access and retention rules alongside the privacy expectations of communities.'),
    ('Better recommendation quality.', 'Evaluate the current rule-based feed with user feedback, improve topic diversity, and give users more control over the interests used to arrange their feed.'),
    ('Additional community controls.', 'Introduce optional membership approval for private communities, delegated community moderation, and more detailed rules for community ownership and management.'),
    ('More complete testing.', 'Add automated integration tests for permissions, OTP limits, comment chains, and moderation records. Conduct mobile, keyboard, and assistive-technology checks with recorded results.'),
    ('Improved media storage.', 'Move uploads to managed object storage where appropriate, add image optimisation and video processing, and plan database and file backups together.'),
    ('Wider device and language support.', 'Consider an installable progressive web application or a mobile application, together with Burmese and English interface options and further accessibility improvements.'),
    ('Deployment and performance evaluation.', 'Measure the behaviour of feeds, uploads, and chat under realistic concurrent use. Use those measurements to guide indexing, caching, monitoring, and hosting decisions.')
]:
    label(title, text)

# Reference-style three-column contents tables, with live Word page references.
def index_table(anchor, headers, entries, is_toc=False):
    t = table(headers, [], [.73, 4.94, .58])
    for entry in entries:
        number, title, target = entry[:3]
        level = entry[3] if is_toc else 0
        cells = t.add_row().cells
        for c, width in zip(cells, [.73, 4.94, .58]):
            c.width = Inches(width)
        cell_text(cells[0], number, bold=level == 1, size=10.5)
        title_p = cell_text(cells[1], ('Chapter ' + number + ': ' if level == 1 else '') + title,
                            bold=level == 1, size=10.5)
        if level == 3:
            title_p.paragraph_format.left_indent = Inches(.15)
        cp = cell_text(cells[2], '')
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field(cp, f'PAGEREF {target} \\h')
        cant = OxmlElement('w:cantSplit')
        t.rows[-1]._tr.get_or_add_trPr().append(cant)
    anchor._p.addnext(t._tbl)

index_table(toc_anchor, ['No.', 'Title', 'Page'], toc, True)
index_table(fig_anchor, ['No.', 'Figure Title', 'Page'], figures)

update = OxmlElement('w:updateFields')
update.set(qn('w:val'), 'true')
doc.settings.element.append(update)
doc.save(OUT)
print(f'Created: {OUT}')
print(f'Headings: {len(toc)}; Figure placeholders: {len(figures)}')
print(f'Body word count: {sum(len(re.findall(r"\S+", p.text)) for p in doc.paragraphs)}')
