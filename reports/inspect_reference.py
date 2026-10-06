from pathlib import Path
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

source = Path(r'C:\Users\Asus\Documents\AzureMed Hub Final 1,1 Report (2) (2) (2).docx')
sys.stdout.reconfigure(encoding='utf-8')
start = int(sys.argv[1]) if len(sys.argv) > 1 else 0
doc = Document(source)
for i, section in enumerate(doc.sections):
    print('SECTION', i, 'page', section.page_width.inches, section.page_height.inches,
          'margins', section.top_margin.inches, section.bottom_margin.inches,
          section.left_margin.inches, section.right_margin.inches)
    print('HEADER', [p.text for p in section.header.paragraphs])
    print('FOOTER', [p.text for p in section.footer.paragraphs])
for name in ['Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'Title']:
    style = doc.styles[name]
    print('STYLE', name, style.font.name, style.font.size.pt if style.font.size else None,
          style.font.bold, 'spacing', style.paragraph_format.line_spacing)
for i, element in enumerate(doc.element.body):
    if i < start:
        continue
    if element.tag == qn('w:p'):
        p = Paragraph(element, doc)
        text = p.text.strip()
        if text:
            fmt = [(r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs if r.text.strip()]
            print(f'P{i} [{p.style.name}] [align={p.alignment}] {text}')
            if len(text) < 160:
                print(' FORMAT', fmt[:5])
        if element.xpath('.//w:drawing'):
            print(f'P{i} [IMAGE]')
        if element.xpath('.//w:br[@w:type="page"]'):
            print(f'P{i} [PAGE BREAK]')
    elif element.tag == qn('w:tbl'):
        table = Table(element, doc)
        print(f'TABLE{i} {len(table.rows)} rows x {len(table.columns)} cols')
        for row in table.rows:
            print(' | '.join(c.text.replace('\n', ' / ') for c in row.cells))
